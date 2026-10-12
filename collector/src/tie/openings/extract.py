"""Text extraction from tender opening-of-offers documents (PDF, scanned PDF, DOCX, legacy
Word 97 DOC via `catdoc`, ODT)."""

import io
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

MIN_TEXT_CHARS = 200
PDFTOTEXT_TIMEOUT = 60
PDFTOPPM_TIMEOUT = 300
TESSERACT_TIMEOUT = 120
CATDOC_TIMEOUT = 60
# OCR renders at most this many pages (`pdftoppm -l`). Opening documents have 1-3 pages; a
# misclassified long scan would otherwise fill /tmp with 300 dpi PNGs and block the single
# worker for hours. Text past the cap is dropped, so such a document fails the cross-check.
OCR_MAX_PAGES = 20
# Pages scanned in landscape come out of `--psm 6` as garbage. Tesseract's orientation
# detection (`--psm 0`, needs the `osd` model) finds them; such pages are read with `--psm 1`
# (automatic segmentation with orientation), the others keep `--psm 6` (one uniform block,
# best for tables). Seen on real scans: upright pages report confidence 5-17, turned 8-11.
OSD_MIN_CONFIDENCE = 2.0
_OSD_ROTATE = re.compile(r"Rotate:\s*(\d+)")
_OSD_CONFIDENCE = re.compile(r"Orientation confidence:\s*([0-9.]+)")

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # Word 97-2003 (.doc) container
_ZIP_MAGIC = b"PK\x03\x04"  # .docx saved under a .doc name
ODT_MAX_XML_BYTES = 50 * 1024 * 1024
_ODF_TEXT = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
_ODF_TABLE = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"


class UnsupportedDocument(Exception):
    """Raised for file extensions that cannot be extracted."""


@dataclass(frozen=True)
class Extracted:
    text: str
    method: str  # "pdftext" | "docx" | "ocr" | "doc" | "odt"


def extract_text(data: bytes, file_name: str) -> Extracted:
    ext = Path(file_name).suffix.lower()
    if ext == ".pdf":
        extracted = _extract_pdf(data)
    elif ext == ".docx":
        extracted = Extracted(text=_extract_docx(data), method="docx")
    elif ext == ".doc":
        extracted = _extract_doc(data, file_name)
    elif ext == ".odt":
        extracted = Extracted(text=_extract_odt(data), method="odt")
    else:
        raise UnsupportedDocument(f"unsupported document type: {file_name!r}")
    return Extracted(text=_normalize_spaces(extracted.text), method=extracted.method)


def _normalize_spaces(text: str) -> str:
    """Amounts use no-break spaces as thousands separators; make them plain spaces."""
    return text.replace(" ", " ").replace(" ", " ")


def _extract_pdf(data: bytes) -> Extracted:
    text = _pdftotext(data)
    if len("".join(text.split())) >= MIN_TEXT_CHARS:
        return Extracted(text=text, method="pdftext")
    return Extracted(text=_ocr(data), method="ocr")


def _pdftotext(data: bytes) -> str:
    try:
        proc = subprocess.run(
            ["pdftotext", "-layout", "-", "-"],
            input=data,
            check=False,
            capture_output=True,
            timeout=PDFTOTEXT_TIMEOUT,
        )
    except subprocess.TimeoutExpired:
        return ""
    if proc.returncode != 0:
        return ""
    return proc.stdout.decode("utf-8", errors="replace")


def _page_number(path: Path) -> int:
    return int(path.stem.rsplit("-", 1)[1])


def _ocr(data: bytes) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        pdf_path = tmp_dir / "file.pdf"
        pdf_path.write_bytes(data)
        subprocess.run(
            ["pdftoppm", "-r", "300", "-l", str(OCR_MAX_PAGES), "-png", str(pdf_path),
             str(tmp_dir / "page")],
            check=True,
            capture_output=True,
            timeout=PDFTOPPM_TIMEOUT,
        )
        pages = sorted(tmp_dir.glob("page-*.png"), key=_page_number)
        texts = []
        for page in pages:
            psm = "1" if _page_rotated(page) else "6"
            proc = subprocess.run(
                ["tesseract", str(page), "-", "-l", "pol", "--psm", psm],
                check=True,
                capture_output=True,
                timeout=TESSERACT_TIMEOUT,
            )
            texts.append(proc.stdout.decode("utf-8", errors="replace"))
    return "\n\f\n".join(texts)


def _page_rotated(page: Path) -> bool:
    """True when tesseract's orientation detection is confident the page is turned.

    Any failure (no `osd` model, too little text on the page, timeout) means "not turned",
    so the page is read as before."""
    try:
        proc = subprocess.run(
            ["tesseract", str(page), "-", "--psm", "0"],
            check=False,
            capture_output=True,
            timeout=TESSERACT_TIMEOUT,
        )
    except subprocess.TimeoutExpired:
        return False
    if proc.returncode != 0:
        return False
    out = proc.stdout.decode("utf-8", errors="replace")
    rotate = _OSD_ROTATE.search(out)
    confidence = _OSD_CONFIDENCE.search(out)
    if not rotate or not confidence:
        return False
    try:
        conf = float(confidence.group(1))
    except ValueError:
        return False
    return int(rotate.group(1)) % 360 != 0 and conf >= OSD_MIN_CONFIDENCE


def _extract_doc(data: bytes, file_name: str) -> Extracted:
    """Legacy Word 97-2003 file. Some ".doc" files are really .docx; those are read as such.

    `catdoc` prints any non-Word input as plain text, so only OLE2 files reach it."""
    if data.startswith(_ZIP_MAGIC):
        return Extracted(text=_extract_docx(data), method="docx")
    if not data.startswith(_OLE2_MAGIC):
        raise UnsupportedDocument(f"not a Word 97-2003 file: {file_name!r}")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "file.doc"
        path.write_bytes(data)
        proc = subprocess.run(
            ["catdoc", "-w", "-d", "utf-8", str(path)],
            check=True,
            capture_output=True,
            timeout=CATDOC_TIMEOUT,
        )
    return Extracted(text=proc.stdout.decode("utf-8", errors="replace"), method="doc")


def _odf_text(el: ElementTree.Element) -> str:
    """Text of an ODF element: spans flattened, <text:s> as spaces, tabs and line breaks
    kept, nested paragraphs (several in one table cell) joined with a space."""
    parts = [el.text or ""]
    for child in el:
        if child.tag == f"{_ODF_TEXT}s":
            parts.append(" " * int(child.get(f"{_ODF_TEXT}c", "1") or "1"))
        elif child.tag == f"{_ODF_TEXT}tab":
            parts.append("\t")
        elif child.tag == f"{_ODF_TEXT}line-break":
            parts.append("\n")
        elif child.tag in (f"{_ODF_TEXT}p", f"{_ODF_TEXT}h"):
            parts.append(" " + _odf_text(child))
        else:
            parts.append(_odf_text(child))
        parts.append(child.tail or "")
    return "".join(parts).strip()


def _extract_odt(data: bytes) -> str:
    """OpenDocument text: paragraphs and headings one per line, table rows as cells joined
    with " | " (the same layout as `_extract_docx`)."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        if z.getinfo("content.xml").file_size > ODT_MAX_XML_BYTES:
            raise UnsupportedDocument("odt content.xml too large")
        root = ElementTree.fromstring(z.read("content.xml"))
    lines: list[str] = []

    def walk(el: ElementTree.Element) -> None:
        for child in el:
            if child.tag == f"{_ODF_TABLE}table-row":
                cells = [c for c in child if c.tag == f"{_ODF_TABLE}table-cell"]
                lines.append(" | ".join(_odf_text(c) for c in cells))
            elif child.tag in (f"{_ODF_TEXT}p", f"{_ODF_TEXT}h"):
                lines.append(_odf_text(child))
            else:
                walk(child)

    walk(root)
    return "\n".join(lines)


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            lines.append(Paragraph(child, doc).text)
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)
